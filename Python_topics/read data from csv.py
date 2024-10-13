import  openpyxl
import pandas as pd
workbook=openpyxl.load_workbook("amazon_cart_items.xlsx")
sheet=workbook["Sheet1"]
row_count=sheet.max_row
column_count=sheet.max_column
print(row_count)


for row in range(1 , row_count+1):
    for col in range(1, column_count+1):
        cell_value=sheet.cell(row=row , column = col).value
        print(cell_value, end = '')
sheet.cell(row= row, column=1)

df=pd.read_csv("odi_most_fifties.xlsx")
print(len(df))


def nested_sum(l):
    total =0
    for i in l:
        if isinstance(i , list):
            total+=nested_sum(i)

        else:
            total +=i
    return total

l =[1, 3 ,4,[4 , 5 ,6] , 89, [3 , 4 ,5] ,6]
# Call the function and print the result
print(nested_sum(l))
